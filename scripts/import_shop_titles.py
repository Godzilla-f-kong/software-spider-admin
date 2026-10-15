"""Import curated software names from the 2026-10-06 shop title workbook.

Run with the project venv. Defaults to a dry run; use --apply to write.
Only `name` and its required normalized form are set, leaving all business
fields at their model defaults. Rows without a clear software product are
intentionally excluded.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import SessionLocal
from app.models.entities import SoftwareCandidate


NAMES = """
2|EssentialPIM Pro
3|GerbView
4|Wise Registry Cleaner Pro
5|PlayerFab
6|KMPlayer
7|SMPlayer
8|BSPlayer
9|Apeaksoft Blu-ray Player
10|Ashampoo Burning Studio
11|Text Edit Plus
12|PDF Annotator
13|Aiseesoft Blu-ray Player
15|TechPowerUp GPU-Z
16|VueScan Pro
17|GPU Caps Viewer
18|QQ Player
19|FxSound
20|Speccy
21|AnyMP4 Blu-ray Player
22|Light Image Resizer
23|Readiris PDF Corporate
24|Iperius Backup
25|Total Uninstall
26|Xilisoft Video Converter Ultimate
27|CyberLink AudioDirector Ultra
29|FastCopy
31|Inpaint
32|Sweet Home 3D
33|Epic Pen Pro
34|PDF-XChange Editor
35|Kodi
36|Tipard FixMP4
37|Wise Data Recovery
38|RS Browser Forensics
39|Ookla Speedtest
40|Axialis IconWorkshop Professional
41|UltraSearch
42|Audacity
43|iMobie DroidKit
44|Shotcut
45|ARCHPR
46|RS RAID Retrieve
47|Calibre
48|IceCream PDF Editor Pro
49|JPEXS Free Flash Decompiler
50|Poweramp
51|Hard Disk Sentinel
52|MediaInfo
54|Room Arranger
55|XnView MP
56|dBpoweramp Music Converter
57|LibreCAD
58|NeeView
59|Advanced Installer Architect
60|PDFMate eBook Converter Professional
62|PrintConductor
63|MobiKin Backup Manager
64|WinTools.net Premium
65|MobiKin Assistant for iOS
66|Affinity Publisher
69|IcoFX
70|Windows Manager
71|Modiac Factory
72|iToolab UnlockGo iPhone
73|UnderCover10
75|LosslessCut
76|EmEditor
77|TelePrompter for Win
79|Aiseesoft iPhone Unlocker
80|CyberLink PowerDVD
81|HWiNFO64
82|PassFab iPhone Unlocker
83|MiniTool Partition Wizard
85|Xournal++
86|PDF24 Creator
90|Rufus
91|BandLab Cakewalk
92|FotoJet Designer
93|iMyFone Passper Pro
94|Ashampoo Backup Pro
95|KeePassXC
96|Freeplane
97|PDF Logo Remover
98|Aiseesoft Video Repair
99|Nitro PDF Pro
100|ShareX
102|FastStone Capture
103|PDFgear
104|Master PDF Editor
105|Telestream Switch Pro
106|Readest
107|iReaShare iPhone Data Eraser
108|Pointerstick
109|Hddb
110|PhotoDirector Ultra
111|WinRAR
112|XYplorer
114|Privacy Eraser Pro
115|PDF Shaper Pro
116|Cisdem PDF Compressor
117|Balabolka
118|SysTools DBX to PST Converter
119|ImgDrive
120|Windows 10 Manager
121|Auslogics BoostSpeed
123|ScreenToGif
124|Mirillis Action
125|Caesium
126|MP3DirectCut
127|PixelOver
129|ASUS GPU Tweak III
131|foobar2000
132|Magic Partition Recovery
133|FFmpeg Batch AV Converter
134|FotoSketcher
136|Tenorshare 4DDiG DLL Fixer
137|SysTools OST Converter
138|Leawo Prof.Media
139|Yuhan Blu-ray Creator
140|ShanaEncoder
141|Stellar Toolkit for File Repair
142|ImageConverter
144|Aiarty Video Enhancer
145|DearMob iPhone Manager
146|HD Video Converter Factory
147|Video Object Remover AI
148|BleachBit
149|MediBang Paint Pro
150|iTop Easy Desktop
151|iToolab UnlockGo Android
152|FixWin
153|MusicBee
154|DBeaver Community
155|4DDiG Duplicate File Deleter
156|RyTuneX
159|Screenpresso Pro
160|CoolUtils Total HTML Converter
161|BitComet
162|Draw.io Desktop
163|PicPick
164|Stellarium
165|Office Tool Plus
166|Picasa
167|FonePaw Screen Recorder
168|MXPlayer
171|AnyBurn
172|DoYourData Super Eraser
173|PowerToys
174|Wipe Pro
175|Inspire Art AI Image Generator
176|Goversoft PrivaZer
177|R-Studio
178|IObit Driver Booster Pro
179|Tenorshare 4DDiG
180|Kdenlive
181|Tenorshare 4uKey for Android
182|Aegisub
183|TunesKit iPhone Unlocker
184|Passper for ZIP
185|CrystalDiskInfo
186|Olib
187|Shutter Encoder
188|Flip PDF Plus Corporate
189|GIMP
190|DDU
191|HeidiSQL
192|Advanced Renamer
194|Librera PRO
195|MusicLyrics
196|Find.Same.Images.OK
197|MKVToolNix
198|4DDiG Partition Manager
199|Equalizer APO
200|LuBan 3D
201|Darktable
202|HyperSnap
203|Avidemux
204|QuickMemoryTestOK
205|IPTV Checker
206|PDF Extra
207|Tenorshare iCareFone
208|010Editor
209|WinToUSB Enterprise
210|qBittorrent
211|LibreOffice
212|AIMP
213|Hasleo Backup Suite
214|Everything
215|Ashampoo Burning Studio
216|Media Player Classic
217|Audio Editor Pro
218|QCAD Professional
219|Beyond Compare
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    sheet = load_workbook(args.workbook, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    if rows[0][0] != "商品标题":
        raise ValueError("Expected 商品标题 in first column")

    names: dict[int, str] = {}
    invalid: list[str] = []
    for line in NAMES.strip().splitlines():
        number, name = line.split("|", 1)
        index = int(number)
        if index in names or index > len(rows):
            raise ValueError(f"Invalid mapped row {index}")
        title = str(rows[index - 1][0] or "")
        # Every mapped name must be anchored in the title. Spaces/case differ.
        if "".join(name.lower().split()) not in "".join(title.lower().split()):
            invalid.append(f"Name {name!r} not found in row {index}: {title}")
        names[index] = name

    if invalid:
        raise ValueError("\n".join(invalid))

    unique: dict[str, str] = {}
    for name in names.values():
        unique.setdefault(name.casefold(), name)

    with SessionLocal() as db:
        existing = set(db.scalars(select(SoftwareCandidate.normalized_name)).all())
        pending = [(name, key) for key, name in unique.items() if key not in existing]
        print(f"titles={len(rows)-1} mapped={len(names)} unique={len(unique)}")
        print(f"unmapped_rows={','.join(map(str, sorted(set(range(2, len(rows)+1))-names.keys())))}")
        print(f"already_present={len(unique)-len(pending)} to_import={len(pending)}")
        if args.apply:
            db.add_all(SoftwareCandidate(name=name, normalized_name=key) for name, key in pending)
            db.commit()
            print(f"imported={len(pending)}")
        else:
            for name, _ in pending:
                print(name)


if __name__ == "__main__":
    main()
